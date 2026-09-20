from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def build_lists_sheet(workbook, enum_lists: dict[str, list[str]]) -> dict[str, str]:
    """Writes one column per enum to a hidden 'Lists' sheet and returns
    {enum_name: excel_range} so dropdown validations can reference it — this
    is the same enum data (app/core/enums.py) the UI dropdowns use, so the
    Excel workbook and the app can never define the allowed values twice.
    """
    ws = workbook.create_sheet("Lists")
    ws.sheet_state = "hidden"

    ranges = {}
    for col_idx, (name, options) in enumerate(enum_lists.items(), start=1):
        letter = get_column_letter(col_idx)
        ws.cell(row=1, column=col_idx, value=name)
        for row_idx, option in enumerate(options, start=2):
            ws.cell(row=row_idx, column=col_idx, value=option)
        ranges[name] = f"'Lists'!${letter}$2:${letter}${len(options) + 1}"
    return ranges


def add_dropdown(ws, column_letter: str, first_row: int, last_row: int, source_range: str) -> None:
    dv = DataValidation(type="list", formula1=source_range, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}{first_row}:{column_letter}{last_row}")
