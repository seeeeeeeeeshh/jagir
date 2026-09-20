import re

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.ui import theme

_INVALID_TABLE_NAME_CHARS = re.compile(r"[^A-Za-z0-9_]")

_BODY_FONT = Font(name="Calibri", size=10.5, color="1A1A1A")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color=theme.PANEL.lstrip("#").upper(), end_color=theme.PANEL.lstrip("#").upper(), fill_type="solid")
_THIN_BORDER = Border(*(Side(style="thin", color="E0E0E0") for _ in range(4)))


def _safe_table_name(name: str) -> str:
    safe = _INVALID_TABLE_NAME_CHARS.sub("", name)
    if not safe or not safe[0].isalpha():
        safe = f"T{safe}"
    return safe


def style_sheet(ws, table_name: str, row_striping: bool = True, freeze_columns: int = 0) -> None:
    """Freeze the header row, turn the sheet into a filterable Excel Table
    (openpyxl requires at least one data row for this), size columns to
    their content, and apply a consistent look: no default gridlines (a
    thin border on every cell reads as designed rather than raw data),
    dark header matching the app's own panel color with bold white text,
    one consistent font throughout. `row_striping=False` for sheets that
    get their own status-driven row tint instead (see
    conditional_format.tint_rows_by_category) — the two looked muddy
    layered together. `freeze_columns` additionally pins that many leading
    columns (e.g. 3 keeps Company/Role in view on a 24-column sheet while
    scrolling right through dates/status/notes)."""
    if ws.max_row < 1 or ws.max_column < 1:
        return

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"{get_column_letter(freeze_columns + 1)}2" if freeze_columns else "A2"

    if ws.max_row >= 2:
        table = Table(displayName=_safe_table_name(table_name), ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=row_striping, showFirstColumn=False
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows():
        for cell in row:
            cell.border = _THIN_BORDER
            if cell.row == 1:
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
                cell.alignment = Alignment(vertical="center")
            else:
                cell.font = _BODY_FONT

    ws.row_dimensions[1].height = 20

    for column_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column_cells)
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 50)


def add_header_comment(ws, df, header: str, text: str) -> None:
    """A note on hover, on the header cell only — never on data cells (that
    gets noisy fast). Used for the Status/Urgency/Application Health
    columns to explain what the color coding means, since a color-only
    signal isn't self-explanatory to someone opening the file without the
    app open next to it."""
    if df.empty or header not in df.columns:
        return
    col_letter = get_column_letter(df.columns.get_loc(header) + 1)
    comment = Comment(text, "Jagir")
    comment.width, comment.height = 260, 100
    ws[f"{col_letter}1"].comment = comment
