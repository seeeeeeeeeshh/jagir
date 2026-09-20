import logging
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from app.core.enums import (
    APPLICATION_HEALTH_OPTIONS,
    EMPLOYMENT_TYPE_OPTIONS,
    INTERVIEW_STAGE_OPTIONS,
    NETWORKING_STATUS_OPTIONS,
    POSITION_TYPE_OPTIONS,
    PREP_CATEGORY_OPTIONS,
    PRIORITY_OPTIONS,
    SOURCE_OPTIONS,
    STATUS_OPTIONS,
    TARGET_REACH_SAFETY_OPTIONS,
    URGENCY_TIERS,
    WORKPLACE_TYPE_OPTIONS,
)
from app.excel.conditional_format import ROW_BUFFER, add_category_conditional_formatting, hyperlink_column
from app.excel.sheet_builders.companies_sheet import build_companies
from app.excel.sheet_builders.dashboard_sheet import build_dashboard_stats, build_status_breakdown
from app.excel.sheet_builders.interviews_sheet import build_interviews
from app.excel.sheet_builders.networking_sheet import build_networking
from app.excel.sheet_builders.opportunities_sheets import (
    build_all_opportunities,
    build_applications,
    build_research,
    build_saved_jobs,
)
from app.excel.sheet_builders.prep_resources_sheet import build_prep_resources
from app.excel.styles import add_header_comment, style_sheet
from app.excel.validation import add_dropdown, build_lists_sheet
from app.ui import theme

logger = logging.getLogger(__name__)

# Same enum data the UI dropdowns use (app/core/enums.py) — the workbook and
# the app can never define the allowed values twice.
ENUM_LISTS = {
    "Status": STATUS_OPTIONS,
    "Priority": PRIORITY_OPTIONS,
    "TargetReachSafety": TARGET_REACH_SAFETY_OPTIONS,
    "PositionType": POSITION_TYPE_OPTIONS,
    "Source": SOURCE_OPTIONS,
    "Workplace": WORKPLACE_TYPE_OPTIONS,
    "EmploymentType": EMPLOYMENT_TYPE_OPTIONS,
    "NetworkingStatus": NETWORKING_STATUS_OPTIONS,
    "InterviewStage": INTERVIEW_STAGE_OPTIONS,
    "PrepCategory": PREP_CATEGORY_OPTIONS,
}

# sheet name -> {column header: enum key in ENUM_LISTS}
DROPDOWN_COLUMNS = {
    "All Opportunities": {
        "Status": "Status", "Priority": "Priority",
        "Target / Reach / Safety": "TargetReachSafety",
        "Position Type": "PositionType", "Workplace": "Workplace",
    },
    "Applications": {
        "Status": "Status", "Priority": "Priority",
        "Target / Reach / Safety": "TargetReachSafety", "Workplace": "Workplace",
    },
    "Research": {
        "Status": "Status", "Priority": "Priority",
        "Target / Reach / Safety": "TargetReachSafety",
    },
    "Saved Jobs": {
        "Status": "Status", "Priority": "Priority",
        "Target / Reach / Safety": "TargetReachSafety",
    },
    "Networking": {"Relationship / Status": "NetworkingStatus"},
    "Interviews": {"Stage": "InterviewStage"},
    "Prep Resources": {"Category": "PrepCategory"},
}

SHEET_BUILDERS = {
    "All Opportunities": build_all_opportunities,
    "Applications": build_applications,
    "Research": build_research,
    "Saved Jobs": build_saved_jobs,
    "Networking": build_networking,
    "Interviews": build_interviews,
    "Prep Resources": build_prep_resources,
    "Companies": build_companies,
}

DASHBOARD_STATUS_START_COL = 3  # 0-indexed pandas startcol -> Excel column D

# Sheets built from the opportunities table (opportunity_rows.py) share the
# same column set, so the same coloring/formula treatment applies to all
# four. Colors reuse app/ui/theme.py's own category->color mapping — the
# same accents the status/urgency/health badges use in the app itself —
# rather than inventing a separate palette for the spreadsheet. Each entry
# is (all possible values, value->color, whole_row) — "all possible
# values", not just what's currently in the sheet, since a live
# Conditional Formatting rule has to cover anything the dropdown could be
# switched to, not just today's data. Urgency/Application Health are listed
# (and therefore added) before Status precisely because they're narrower
# (single-column) rules that need higher CF priority than Status's
# whole-row rule to win at their own cells — see
# conditional_format.add_category_conditional_formatting's docstring.
_COLOR_COLUMNS: list[tuple[str, list[str], callable, bool]] = [
    ("Urgency", URGENCY_TIERS, lambda v: theme.category_color(theme.urgency_category(v)), False),
    ("Application Health", APPLICATION_HEALTH_OPTIONS, lambda v: theme.category_color(theme.health_category(v)), False),
    ("Status", STATUS_OPTIONS, lambda v: theme.category_color(theme.status_category(v)), True),
]
_OPPORTUNITY_SHEETS = {"All Opportunities", "Applications", "Research", "Saved Jobs"}
# Colors alone aren't self-explanatory to someone opening the file without
# the app next to it — these land as a hover comment on each header cell.
_COLOR_LEGENDS = {
    "Status": (
        "Row color follows pipeline stage — and updates live if you change "
        "Status yourself via the dropdown:\n"
        "gray = not yet applied\n"
        "blue, lighter -> darker = applied through interviews (darker = further along)\n"
        "amber = on hold\n"
        "green, lighter -> darker = offer -> accepted\n"
        "rose = rejected\n"
        "gray = withdrawn / ghosted"
    ),
    "Urgency": "amber = deadline soon or critical\ngray = normal or no deadline",
    "Application Health": "amber = needs action or gone stale\ngray = active, nothing needed right now",
}
# Application ID/Company/Role stay in view while scrolling right through
# the other 21 columns — a 24-wide sheet is unreadable without this.
_FREEZE_COLUMNS = 3


def export_workbook(conn: sqlite3.Connection, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    frames = {name: builder(conn) for name, builder in SHEET_BUILDERS.items()}
    dashboard_stats = build_dashboard_stats(conn)
    status_breakdown = build_status_breakdown(conn)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, df in frames.items():
            df.to_excel(writer, sheet_name=name, index=False)

        dashboard_stats.to_excel(writer, sheet_name="Dashboard", index=False, startrow=0, startcol=0)
        status_breakdown.to_excel(
            writer, sheet_name="Dashboard", index=False, startrow=0, startcol=DASHBOARD_STATUS_START_COL
        )

        workbook = writer.book
        list_ranges = build_lists_sheet(workbook, ENUM_LISTS)

        for name, df in frames.items():
            ws = writer.sheets[name]
            # The opportunity sheets get a status-driven row tint instead of
            # plain zebra-striping — the two layered together read as muddy,
            # so table-level banding is off wherever the tint is on.
            style_sheet(
                ws, table_name=name,
                row_striping=name not in _OPPORTUNITY_SHEETS,
                freeze_columns=_FREEZE_COLUMNS if name in _OPPORTUNITY_SHEETS else 0,
            )
            for header, enum_key in DROPDOWN_COLUMNS.get(name, {}).items():
                if df.empty or header not in df.columns:
                    continue
                col_letter = get_column_letter(df.columns.get_loc(header) + 1)
                add_dropdown(ws, col_letter, 2, len(df) + 1 + ROW_BUFFER, list_ranges[enum_key])

            if name in _OPPORTUNITY_SHEETS:
                for header, possible_values, color_for_value, whole_row in _COLOR_COLUMNS:
                    add_category_conditional_formatting(
                        ws, df, header, possible_values, color_for_value, whole_row=whole_row,
                        opacity=0.30 if whole_row else 1.0,
                    )
                for header, legend in _COLOR_LEGENDS.items():
                    add_header_comment(ws, df, header, legend)
                hyperlink_column(ws, df, "Job URL")
                _write_days_since_application_formula(ws, df)

            if name == "Networking":
                hyperlink_column(ws, df, "LinkedIn URL")

        dashboard_ws = writer.sheets["Dashboard"]
        dashboard_ws.freeze_panes = "A2"
        for col_letter, width in (("A", 30), ("B", 12), ("D", 20), ("E", 10)):
            dashboard_ws.column_dimensions[col_letter].width = width
        _add_status_chart(dashboard_ws, status_breakdown)

    logger.info("Exported workbook to %s (%d opportunities)", output_path, len(frames["All Opportunities"]))
    return output_path


def _write_days_since_application_formula(ws, df: pd.DataFrame) -> None:
    """Live Excel formula instead of a value computed at export time — the
    number updates every time the sheet is opened, with no re-export
    needed, and stays blank for anything not yet applied to. Extends past
    the exported rows by the same buffer as the conditional formatting and
    dropdowns, so a row you type in yourself still gets a working formula,
    not a dead blank cell."""
    if df.empty or "Date Applied" not in df.columns or "Days Since Application" not in df.columns:
        return
    applied_col = get_column_letter(df.columns.get_loc("Date Applied") + 1)
    days_col = get_column_letter(df.columns.get_loc("Days Since Application") + 1)
    for row_idx in range(2, len(df) + 2 + ROW_BUFFER):
        ref = f"{applied_col}{row_idx}"
        cell = ws[f"{days_col}{row_idx}"]
        cell.value = f'=IF(ISBLANK({ref}),"",TODAY()-{ref})'
        cell.number_format = "0"
    ws.column_dimensions[days_col].width = max(ws.column_dimensions[days_col].width or 0, 12)


def _add_status_chart(ws, status_breakdown: pd.DataFrame) -> None:
    if status_breakdown.empty:
        return
    n = len(status_breakdown)
    cat_col = DASHBOARD_STATUS_START_COL + 1  # 1-indexed "Status" column
    val_col = DASHBOARD_STATUS_START_COL + 2  # 1-indexed "Count" column

    chart = BarChart()
    chart.title = "Opportunities by status"
    chart.y_axis.title = "Count"
    data = Reference(ws, min_col=val_col, min_row=1, max_row=n + 1)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=n + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "A10")
