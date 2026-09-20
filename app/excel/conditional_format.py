"""Colors a column's cells by their own value using real Excel Conditional
Formatting rules — not a fill baked into the cell once at export time.

That distinction matters: a plain PatternFill is dead the moment you change
the cell's value in Excel (the color just stays whatever it was when the
file was generated). A Conditional Formatting rule is a live formula Excel
re-evaluates on every recalculation, so switching an application's Status
via its own dropdown updates the color immediately, in Excel itself, no
re-export needed. Colors reuse app/ui/theme.py's own category->color
mapping — the same accents the app's own badges use — so the exported
sheet never invents a separate palette.
"""

from __future__ import annotations

from typing import Callable

import pandas as pd
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Manually-added rows past the exported data still get live coloring, up to
# this many extra rows — covers ordinary manual additions without bloating
# the file out to Excel's full 1,048,576-row ceiling for no reason. Shared
# with export.py's dropdown validation and Days-Since-Application formula,
# which extend by the same amount for the same reason.
ROW_BUFFER = 300


def add_category_conditional_formatting(
    ws,
    df: pd.DataFrame,
    header: str,
    possible_values: list[str],
    color_for_value: Callable[[str], str | None],
    *,
    whole_row: bool = False,
    opacity: float = 1.0,
) -> None:
    """One CF rule per possible value (not just values currently present in
    `df` — the whole point is that switching a cell to any valid option
    colors correctly, including options nothing in the sheet uses yet).

    `whole_row=True` colors every column in that row (Status); `False`
    colors only this column (Urgency, Application Health). Call the
    column-scoped ones *before* the whole-row one for the same overlapping
    range — Conditional Formatting Rules that openpyxl adds first get a
    lower priority number, which Excel evaluates first, and stopIfTrue=True
    there keeps the generic row rule from also painting over that same
    cell.
    """
    if df.empty or header not in df.columns:
        return
    col_letter = get_column_letter(df.columns.get_loc(header) + 1)
    last_row = len(df) + 1 + ROW_BUFFER

    if whole_row:
        last_col_letter = get_column_letter(len(df.columns))
        cell_range = f"A2:{last_col_letter}{last_row}"
    else:
        cell_range = f"{col_letter}2:{col_letter}{last_row}"

    for value in possible_values:
        color = color_for_value(value)
        if not color:
            continue
        fill_hex = _blend_with_white(color, opacity) if opacity < 1.0 else color.lstrip("#").upper()
        fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")

        if whole_row:
            # $-anchored column, relative row — Excel re-targets the row
            # part per row in the range while always reading the same
            # (this row's) Status cell.
            rule = FormulaRule(formula=[f'${col_letter}2="{value}"'], fill=fill)
        else:
            rule = CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill, stopIfTrue=True)
        ws.conditional_formatting.add(cell_range, rule)


def hyperlink_column(ws, df: pd.DataFrame, header: str, link_color: str = "1155CC") -> None:
    """A URL column is otherwise just text you have to copy-paste — Excel
    doesn't auto-linkify a plain string cell the way it does when you type
    a URL by hand. Sets a real clickable Hyperlink plus link-styled text
    (style_sheet's uniform body font runs first, so this needs to override
    it explicitly here, not just set .hyperlink alone)."""
    if df.empty or header not in df.columns:
        return
    col_letter = get_column_letter(df.columns.get_loc(header) + 1)
    for row_idx, value in enumerate(df[header], start=2):
        if not value:
            continue
        cell = ws[f"{col_letter}{row_idx}"]
        cell.hyperlink = value
        cell.font = Font(name=cell.font.name, size=cell.font.size, color=link_color, underline="single")


def _blend_with_white(hex_color: str, opacity: float) -> str:
    hex_color = hex_color.lstrip("#")
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    blended = (round(c * opacity + 255 * (1 - opacity)) for c in (r, g, b))
    return "".join(f"{c:02X}" for c in blended)
