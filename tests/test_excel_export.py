import openpyxl

from app.core.models import Contact, Interview, Opportunity, PrepResource
from app.db.repositories.companies_repo import CompaniesRepo
from app.db.repositories.contacts_repo import ContactsRepo
from app.db.repositories.interviews_repo import InterviewsRepo
from app.db.repositories.opportunities_repo import OpportunitiesRepo
from app.db.repositories.prep_resources_repo import PrepResourcesRepo
from app.excel.export import SHEET_BUILDERS, export_workbook

EXPECTED_SHEETS = {
    "All Opportunities", "Applications", "Research", "Saved Jobs",
    "Networking", "Interviews", "Prep Resources", "Companies", "Dashboard",
}


def _seed(conn):
    companies = CompaniesRepo(conn)
    opportunities = OpportunitiesRepo(conn)
    contacts = ContactsRepo(conn)
    interviews = InterviewsRepo(conn)
    prep = PrepResourcesRepo(conn)

    acme = companies.get_or_create("Acme Corp")
    lab = companies.get_or_create("Some University Lab")

    applied = opportunities.create(
        Opportunity(
            title="Backend Intern", company_id=acme.id, position_type="Internship",
            status="Applied", source="linkedin", origin="linkedin_import",
            date_saved="2026-01-01", date_applied="2026-01-05",
        )
    )
    opportunities.create(
        Opportunity(
            title="RA — Vision Lab", company_id=lab.id, position_type="RA",
            status="Saved", source="university", origin="manual", date_saved="2026-01-02",
        )
    )

    contact = contacts.create(Contact(name="Jane Recruiter", company_id=acme.id, relationship="Connected"))
    contacts.link_to_opportunity(applied.id, contact.id, role="recruiter")

    interviews.create(Interview(opportunity_id=applied.id, stage="OA", scheduled_at="2026-01-10"))

    prep.create(PrepResource(title="LeetCode Arrays", category="LeetCode"))


def test_export_produces_all_expected_sheets(conn, tmp_path):
    _seed(conn)
    output_path = tmp_path / "export.xlsx"

    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    assert EXPECTED_SHEETS.issubset(set(wb.sheetnames))
    assert "Lists" in wb.sheetnames
    assert wb["Lists"].sheet_state == "hidden"


def test_applications_sheet_has_frozen_header_and_expected_row(conn, tmp_path):
    _seed(conn)
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Applications"]
    # Header row + first 3 columns (Application ID/Company/Role) frozen —
    # a 24-column sheet is unreadable without pinning identity columns
    # while scrolling right.
    assert ws.freeze_panes == "D2"

    header = [c.value for c in ws[1]]
    assert "Role" in header
    assert "Company" in header

    role_col = header.index("Role") + 1
    company_col = header.index("Company") + 1
    assert ws.cell(row=2, column=role_col).value == "Backend Intern"
    assert ws.cell(row=2, column=company_col).value == "Acme Corp"


def test_saved_jobs_sheet_excludes_applied_records(conn, tmp_path):
    _seed(conn)
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Saved Jobs"]
    header = [c.value for c in ws[1]]
    role_col = header.index("Role") + 1
    roles = [ws.cell(row=r, column=role_col).value for r in range(2, ws.max_row + 1)]

    assert "RA — Vision Lab" in roles
    assert "Backend Intern" not in roles  # already applied -> lives in Applications, not here


def test_export_on_empty_database_does_not_error(conn, tmp_path):
    output_path = tmp_path / "empty_export.xlsx"
    path = export_workbook(conn, output_path)
    assert path.exists()

    wb = openpyxl.load_workbook(path)
    assert EXPECTED_SHEETS.issubset(set(wb.sheetnames))


def test_sheet_builders_cover_every_non_dashboard_sheet():
    assert set(SHEET_BUILDERS.keys()) == EXPECTED_SHEETS - {"Dashboard"}


def test_companies_sheet_counts_applications_by_status_not_date_applied(conn, tmp_path):
    """Regression test: date_applied is frequently never captured (LinkedIn
    list-page scrapes don't expose it) — Status reliably is, so counting
    "applications" by date_applied being set previously meant this column
    silently read 0 for every company regardless of real applications."""
    companies = CompaniesRepo(conn)
    opportunities = OpportunitiesRepo(conn)
    acme = companies.get_or_create("Acme Corp")
    opportunities.create(Opportunity(title="Saved role", company_id=acme.id, position_type="Internship", status="Saved", source="linkedin", date_applied=None))
    opportunities.create(Opportunity(title="Applied role", company_id=acme.id, position_type="Internship", status="Applied", source="linkedin", date_applied=None))

    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Companies"]
    header = [c.value for c in ws[1]]
    applications_col = header.index("Applications") + 1
    assert ws.cell(row=2, column=applications_col).value == 1


def test_job_url_is_a_real_clickable_hyperlink(conn, tmp_path):
    companies = CompaniesRepo(conn)
    opportunities = OpportunitiesRepo(conn)
    acme = companies.get_or_create("Acme Corp")
    opportunities.create(
        Opportunity(
            title="Backend Intern", company_id=acme.id, position_type="Internship",
            status="Applied", source="linkedin",
            job_url="https://example.com/jobs/123",
        )
    )
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["All Opportunities"]
    header = [c.value for c in ws[1]]
    job_url_col = header.index("Job URL") + 1
    cell = ws.cell(row=2, column=job_url_col)

    assert cell.hyperlink is not None
    assert cell.hyperlink.target == "https://example.com/jobs/123"
    assert cell.font.underline == "single"


def test_networking_linkedin_url_is_a_real_clickable_hyperlink(conn, tmp_path):
    contacts = ContactsRepo(conn)
    contacts.create(Contact(name="Jane Recruiter", linkedin_url="https://linkedin.com/in/jane"))

    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Networking"]
    header = [c.value for c in ws[1]]
    url_col = header.index("LinkedIn URL") + 1
    cell = ws.cell(row=2, column=url_col)

    assert cell.hyperlink is not None
    assert cell.hyperlink.target == "https://linkedin.com/in/jane"


def test_status_coloring_is_live_conditional_formatting_not_a_static_fill(conn, tmp_path):
    """The bug this locks in: a static PatternFill baked in at export time
    never changes when the user edits the cell in Excel afterward. A real
    Conditional Formatting rule does — Excel re-evaluates it continuously."""
    _seed(conn)
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["All Opportunities"]
    header = [c.value for c in ws[1]]
    status_col = header.index("Status") + 1

    # No static fill baked into the cell itself...
    assert ws.cell(row=2, column=status_col).fill.fill_type is None

    # ...instead, a real rule set covering the whole row, keyed off Status.
    ranges = [cf.sqref for cf in ws.conditional_formatting]
    assert any(str(r).startswith("A2:") for r in ranges)


def test_status_conditional_formatting_covers_every_possible_status_not_just_seeded_ones(conn, tmp_path):
    """The whole point: switching a cell to ANY valid Status (not just the
    two present in this seed data) must color correctly, because the CF
    rule set has to cover the full enum, not just today's data."""
    from app.core.enums import STATUS_OPTIONS

    _seed(conn)
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["All Opportunities"]
    formulas_covered = set()
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.formula:
                formulas_covered.add(rule.formula[0])

    for status in STATUS_OPTIONS:
        assert any(f'"{status}"' in f for f in formulas_covered), f"no CF rule covers status {status!r}"


def test_urgency_and_health_rules_outrank_status_row_rule_at_their_own_cells(conn, tmp_path):
    """Urgency/Application Health need their own color even inside a row
    that's tinted for Status — achieved via CF priority (lower number wins)
    plus stopIfTrue on the narrower rules. If this regresses, the whole-row
    Status tint would silently paint over the Urgency/Health signal."""
    _seed(conn)
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["All Opportunities"]

    priorities_by_range = {}
    for cf in ws.conditional_formatting:
        range_str = str(cf.sqref)
        priorities = [r.priority for r in cf.rules]
        priorities_by_range[range_str] = priorities
        if range_str.startswith("A2:"):
            assert all(r.stopIfTrue is not True for r in cf.rules), "Status's whole-row rules should not stop lower rules"

    status_row_priorities = next(p for r, p in priorities_by_range.items() if r.startswith("A2:"))
    other_priorities = [p for r, ps in priorities_by_range.items() if not r.startswith("A2:") for p in ps]

    assert max(other_priorities) < min(status_row_priorities), (
        "Urgency/Application Health rules must have lower priority numbers (higher precedence) than Status's row rule"
    )


def test_status_header_has_a_color_legend_comment(conn, tmp_path):
    _seed(conn)
    output_path = tmp_path / "export.xlsx"
    export_workbook(conn, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["All Opportunities"]
    header = [c.value for c in ws[1]]
    status_col = header.index("Status") + 1
    comment = ws.cell(row=1, column=status_col).comment

    assert comment is not None
    assert "pipeline stage" in comment.text
